"""Tests for app.ingest.

Unit tests use plain parsed rows / xlsx files and need no database (rows that
are rejected before any entity is touched never use the session). Integration
tests use the rolled-back `db_session` fixture and small xlsx files built with
openpyxl in tmp_path.
"""

from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select

from app.ingest import (
    LOGGER_ALIASES,
    TEMP_ALIASES,
    TIME_ALIASES,
    IngestError,
    ParsedRow,
    canonical_header,
    ingest_rows,
    ingest_xlsx,
    raw_text,
    read_xlsx_rows,
    resolve_sheet_layout,
)
from app.models import Fridge, Logger, LoggerAssignment, TemperatureReading
from app.normalize import parse_timestamp

COMBINED = ["Logger", "Branch", "Fridge", "Time", "Temp"]


def make_xlsx(tmp_path, header, rows, name="file.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    if header is not None:
        ws.append(header)
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


def prow(number, logger="TL-1", branch="Tel Aviv", fridge="Walk-in", time="2026-09-14 06:00", temp=4.0):
    return ParsedRow(number, logger, branch, fridge, time, temp)


def counts(session):
    return {
        model.__tablename__: session.scalar(select(func.count()).select_from(model))
        for model in (Fridge, Logger, LoggerAssignment, TemperatureReading)
    }


def readings(session):
    return session.scalars(
        select(TemperatureReading).order_by(TemperatureReading.logger_id, TemperatureReading.measured_at)
    ).all()


# ===========================================================================
# Unit tests (no database)
# ===========================================================================


def test_raw_text():
    assert raw_text(38.3) == "38.3"
    assert raw_text(39) == "39"
    assert raw_text("ERR") == "ERR"
    assert raw_text(None) is None
    assert raw_text(datetime(2026, 9, 14, 6, 0)) == "2026-09-14T06:00:00"


def test_blank_row_detection():
    assert ParsedRow(2, None, "  ", "", None, None).is_blank
    assert not prow(2, time=None, temp=None).is_blank


def test_rejected_rows_report_row_number_and_reason_without_database():
    rows = [
        prow(2, time="not a time"),
        prow(3, time=None),
        prow(4, logger=None),
        prow(5, branch="  ", fridge=None),
        ParsedRow(6, None, None, None, None, None),  # blank
        ParsedRow(7, "TL-1", None, None, None, None, error="custom reason"),
    ]
    result = ingest_rows(None, rows, "f.xlsx")
    reasons = {r.row_number: r.reason for r in result.rejected}
    assert "'not a time'" in reasons[2]
    assert "missing" in reasons[3]
    assert reasons[4] == "missing logger"
    assert reasons[5] == "missing branch, fridge"
    assert reasons[7] == "custom reason"
    assert 6 not in reasons
    assert (result.rows_read, result.blank_rows, result.inserted) == (5, 1, 0)


def test_invalid_logger_unit_mapping_is_rejected():
    with pytest.raises(ValueError):
        ingest_rows(None, [], "f.xlsx", logger_units={"TL-1": "K"})


def test_invalid_source_unit_mapping_is_rejected():
    with pytest.raises(ValueError):
        ingest_rows(None, [], "f.xlsx", source_units={"TL-1": "K"})


def test_reader_combined_format_row_numbers_and_raw_cells(tmp_path):
    path = make_xlsx(tmp_path, ["logger", "BRANCH", " Fridge ", "time", "TEMP"], [["TL-1", "Haifa", "Dairy", "14/09/2026 06:00", 38.3]])
    (row,) = read_xlsx_rows(path)
    assert (row.row_number, row.logger_code, row.branch_name, row.fridge_name) == (2, "TL-1", "Haifa", "Dairy")
    assert (row.raw_time, row.raw_temp) == ("14/09/2026 06:00", 38.3)


def test_reader_logger_only_file_uses_explicit_logger_and_fridge(tmp_path):
    path = make_xlsx(tmp_path, ["Timestamp", "Temperature"], [["2026-09-14 06:00", 4.1], [None, None], ["2026-09-14 06:15", 4.2]])
    rows = read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))
    assert [(r.row_number, r.logger_code, r.branch_name, r.fridge_name) for r in rows[:1]] == [
        (2, "TL-9", "Haifa", "Dairy")
    ]
    assert rows[1].is_blank and rows[1].row_number == 3
    assert rows[2].row_number == 4 and rows[2].raw_time == "2026-09-14 06:15"


def test_reader_logger_only_file_without_logger_and_fridge_is_an_error(tmp_path):
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.1]])
    with pytest.raises(IngestError, match="pass the logger and fridge explicitly"):
        read_xlsx_rows(path)


def test_reader_requires_both_logger_and_fridge_when_explicit(tmp_path):
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.1]])
    with pytest.raises(IngestError, match="both logger and fridge"):
        read_xlsx_rows(path, logger="TL-9")
    with pytest.raises(IngestError, match="both logger and fridge"):
        read_xlsx_rows(path, fridge=Fridge(branch_name="Haifa", name="Dairy"))


@pytest.mark.parametrize("header", [["Logger", "Branch", "Fridge", "Temp"], ["Time", "Value"], ["Time"]])
def test_reader_missing_time_or_temp_columns_is_a_file_error(tmp_path, header):
    path = make_xlsx(tmp_path, header, [])
    with pytest.raises(IngestError, match="missing required column"):
        read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))


def test_reader_empty_and_corrupt_files(tmp_path):
    empty = make_xlsx(tmp_path, None, [], "empty.xlsx")
    with pytest.raises(IngestError, match="empty"):
        read_xlsx_rows(empty)
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a zip")
    with pytest.raises(IngestError, match="cannot read"):
        read_xlsx_rows(bad)


def test_reader_flags_rows_for_a_different_logger(tmp_path):
    path = make_xlsx(
        tmp_path,
        ["Logger", "Time", "Temp"],
        [["tl-9", "2026-09-14 06:00", 4.0], ["TL-OTHER", "2026-09-14 06:15", 4.0]],
    )
    rows = read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))
    assert rows[0].error is None and rows[0].fridge_name == "Dairy"
    assert "TL-OTHER" in rows[1].error


def test_reader_fridge_by_id_needs_a_session(tmp_path):
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.0]])
    with pytest.raises(IngestError, match="session"):
        read_xlsx_rows(path, logger="TL-9", fridge=1)


def _explicit_rows(tmp_path, header, rows=None):
    path = make_xlsx(tmp_path, header, rows or [["2026-09-14 06:00", 4.1]])
    return read_xlsx_rows(
        path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy")
    )


@pytest.mark.parametrize("time_header", TIME_ALIASES)
def test_each_time_alias_is_accepted(tmp_path, time_header):
    rows = _explicit_rows(tmp_path, [time_header, "Temp"])
    assert rows[0].raw_time == "2026-09-14 06:00"
    assert rows[0].raw_temp == 4.1


@pytest.mark.parametrize("temp_header, unit", TEMP_ALIASES)
def test_each_temp_alias_is_accepted(tmp_path, temp_header, unit):
    rows = _explicit_rows(tmp_path, ["Time", temp_header])
    assert rows[0].raw_temp == 4.1
    layout = resolve_sheet_layout(["Time", temp_header])
    assert layout.temp_unit == unit


@pytest.mark.parametrize("logger_header", LOGGER_ALIASES)
def test_each_logger_alias_is_accepted(tmp_path, logger_header):
    path = make_xlsx(
        tmp_path,
        [logger_header, "Branch", "Fridge", "Time", "Temp"],
        [["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 4.1]],
    )
    (row,) = read_xlsx_rows(path)
    assert row.logger_code == "TL-1"
    assert row.fridge_name == "Dairy"


def test_reversed_column_order_is_accepted(tmp_path):
    rows = _explicit_rows(tmp_path, ["Temp", "Time"], [[4.2, "2026-09-14 06:15"]])
    assert rows[0].raw_time == "2026-09-14 06:15"
    assert rows[0].raw_temp == 4.2


def test_header_case_whitespace_and_punctuation_variants(tmp_path):
    rows = _explicit_rows(tmp_path, [" DATE / TIME ", " temperature (°c) "])
    assert rows[0].raw_temp == 4.1
    rows = _explicit_rows(tmp_path, ["date-time", "Temp.C"])
    assert rows[0].raw_temp == 4.1
    path = make_xlsx(
        tmp_path,
        ["device-id", "TIME", "TEMP_F"],
        [["TL-9", "2026-09-14 06:00", 4.1]],
    )
    rows = read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))
    assert rows[0].logger_code == "TL-9"
    assert canonical_header("Temperature(°C)") == canonical_header("temperature c")


def test_value_is_not_guessed_as_temperature(tmp_path):
    path = make_xlsx(tmp_path, ["Time", "Value"], [["2026-09-14 06:00", 4.1]])
    with pytest.raises(IngestError, match="missing required column"):
        read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))


def test_ambiguous_time_or_temp_headers_are_rejected(tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    with pytest.raises(IngestError, match="ambiguous column"):
        read_xlsx_rows(
            make_xlsx(tmp_path, ["Time", "Timestamp", "Temp"], []),
            logger="TL-9",
            fridge=fridge,
        )
    with pytest.raises(IngestError, match="ambiguous column"):
        read_xlsx_rows(
            make_xlsx(tmp_path, ["Time", "Temp", "Temperature"], []),
            logger="TL-9",
            fridge=fridge,
        )
    with pytest.raises(IngestError, match="ambiguous column"):
        read_xlsx_rows(
            make_xlsx(tmp_path, ["Logger", "Logger ID", "Time", "Temp"], []),
            logger="TL-9",
            fridge=fridge,
        )


def test_xlsx_date_formatted_cell_is_not_treated_as_midnight(tmp_path):
    """Excel date-only cells become datetime(…, 00:00) in openpyxl; that is not a reading time."""
    path = tmp_path / "excel_date.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Temp"])
    date_cell = ws.cell(row=2, column=1, value=date(2026, 9, 14))
    date_cell.number_format = "YYYY-MM-DD"
    ws.cell(row=2, column=2, value=4.1)
    midnight = ws.cell(row=3, column=1, value=datetime(2026, 9, 14, 0, 0))
    midnight.number_format = "YYYY-MM-DD HH:MM"
    ws.cell(row=3, column=2, value=4.2)
    wb.save(path)

    fridge = Fridge(branch_name="Haifa", name="Dairy")
    rows = read_xlsx_rows(path, logger="TL-9", fridge=fridge)
    assert isinstance(rows[0].raw_time, date) and not isinstance(rows[0].raw_time, datetime)
    assert isinstance(rows[1].raw_time, datetime)
    assert rows[1].raw_time == datetime(2026, 9, 14, 0, 0)

    result = ingest_rows(None, rows[:1], "excel_date.xlsx")
    assert result.inserted == 0
    assert result.rejected[0].row_number == 2
    assert "time of day" in result.rejected[0].reason
    assert parse_timestamp(rows[1].raw_time) == datetime(2026, 9, 14, 0, 0)


def test_date_only_values_fail_timestamp_validation(tmp_path):
    path = make_xlsx(
        tmp_path,
        ["Date", "Temp"],
        [["2026-09-14", 4.1], ["14/09/2026", 4.2]],
    )
    rows = read_xlsx_rows(path, logger="TL-9", fridge=Fridge(branch_name="Haifa", name="Dairy"))
    result = ingest_rows(None, rows, "date.xlsx")
    assert result.inserted == 0
    assert {r.row_number for r in result.rejected} == {2, 3}
    assert all("timestamp" in r.reason.lower() or "time" in r.reason.lower() for r in result.rejected)

    date_row = ParsedRow(2, "TL-9", "Haifa", "Dairy", date(2026, 9, 14), 4.1)
    result = ingest_rows(None, [date_row], "date.xlsx")
    assert result.inserted == 0
    assert "time of day" in result.rejected[0].reason


# ===========================================================================
# Integration tests (real PostgreSQL, rolled back per test)
# ===========================================================================

integration = pytest.mark.integration


@integration
def test_combined_format_end_to_end(db_session, tmp_path):
    path = make_xlsx(
        tmp_path,
        COMBINED,
        [
            ["TL-1", "Tel Aviv", "Walk-in", "2026-09-14 06:00", 4.1],
            ["TL-1", "Tel Aviv", "Walk-in", "2026-09-14 06:15", 9.4],
            ["TL-2", "Haifa", "Dairy", "14/09/2026 06:00", 3.8],
        ],
        "logs_batch.xlsx",
    )
    result = ingest_xlsx(db_session, path)
    assert (result.rows_read, result.inserted, result.duplicates) == (3, 3, 0)
    assert result.conflicts == [] and result.rejected == []
    assert (result.loggers_created, result.fridges_created, result.assignments_created) == (2, 2, 2)
    assert counts(db_session) == {"fridges": 2, "loggers": 2, "logger_assignments": 2, "temperature_readings": 3}
    assert {r.temperature_c for r in readings(db_session)} == {Decimal("4.10"), Decimal("9.40"), Decimal("3.80")}
    assert db_session.scalars(select(Logger.unit)).all() == ["C", "C"]


@integration
def test_logger_only_file_with_explicit_logger_and_fridge_objects(db_session, tmp_path):
    logger = Logger(code="TL-9", unit="C")
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add_all([logger, fridge])
    db_session.flush()
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.1], ["2026-09-14 06:15", "ERR"]])

    result = ingest_xlsx(db_session, path, logger=logger, fridge=fridge)

    assert (result.inserted, result.null_temperatures, result.loggers_created, result.fridges_created) == (2, 1, 0, 0)
    rows = readings(db_session)
    assert [(r.logger_id, r.fridge_id) for r in rows] == [(logger.id, fridge.id)] * 2
    assert result.assignments_created == 1
    assert counts(db_session)["logger_assignments"] == 1


@integration
def test_logger_only_file_with_identifiers(db_session, tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add(fridge)
    db_session.flush()
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.1]])

    result = ingest_xlsx(db_session, path, logger="TL-NEW", fridge=fridge.id, logger_units={"TL-NEW": "F"})

    assert result.inserted == 1 and result.loggers_created == 1
    logger = db_session.scalars(select(Logger)).one()
    assert (logger.code, logger.unit) == ("TL-NEW", "F")
    assert readings(db_session)[0].temperature_c == Decimal("-15.50")  # 4.1 F
    with pytest.raises(IngestError, match="does not exist"):
        ingest_xlsx(db_session, path, logger="TL-NEW", fridge=999999)


@integration
def test_logger_only_file_without_explicit_info_errors_and_writes_nothing(db_session, tmp_path):
    path = make_xlsx(tmp_path, ["Time", "Temp"], [["2026-09-14 06:00", 4.1]])
    with pytest.raises(IngestError):
        ingest_xlsx(db_session, path)
    assert counts(db_session) == {"fridges": 0, "loggers": 0, "logger_assignments": 0, "temperature_readings": 0}


@integration
def test_logger_column_for_another_logger_is_rejected_in_explicit_mode(db_session, tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add(fridge)
    db_session.flush()
    path = make_xlsx(
        tmp_path,
        ["Logger", "Time", "Temp"],
        [["TL-9", "2026-09-14 06:00", 4.0], ["TL-OTHER", "2026-09-14 06:15", 4.0]],
    )
    result = ingest_xlsx(db_session, path, logger="TL-9", fridge=fridge)
    assert result.inserted == 1
    assert [r.row_number for r in result.rejected] == [3]
    assert "TL-OTHER" in result.rejected[0].reason
    assert db_session.scalars(select(Logger.code)).all() == ["TL-9"]


@integration
def test_raw_values_and_source_are_preserved(db_session, tmp_path):
    path = make_xlsx(
        tmp_path,
        COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "14/09/2026 06:00", 38],
            ["TL-1", "Haifa", "Dairy", datetime(2026, 9, 14, 6, 15), 38.3],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:30", "ERR"],
        ],
        "TL-1 export.xlsx",
    )
    ingest_xlsx(db_session, path, logger_units={"TL-1": "F"})
    first, second, third = readings(db_session)
    assert (first.raw_timestamp, first.raw_value, first.source_file, first.source_row) == (
        "14/09/2026 06:00", "38", "TL-1 export.xlsx", 2,
    )
    assert first.measured_at == datetime(2026, 9, 14, 6, 0)
    assert (second.raw_timestamp, second.raw_value, second.source_row) == ("2026-09-14T06:15:00", "38.3", 3)
    assert second.measured_at == datetime(2026, 9, 14, 6, 15)
    assert (third.raw_value, third.source_row) == ("ERR", 4)


@integration
def test_err_is_stored_as_null_with_raw_value(db_session, tmp_path):
    path = make_xlsx(
        tmp_path,
        COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", "ERR"],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", None],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:30", 150],
        ],
    )
    result = ingest_xlsx(db_session, path)
    assert (result.inserted, result.null_temperatures, result.rejected) == (3, 2, [])
    err, blank, hot = readings(db_session)
    assert (err.temperature_c, err.raw_value) == (None, "ERR")
    assert (blank.temperature_c, blank.raw_value) == (None, None)
    assert hot.temperature_c == Decimal("150.00")  # no threshold rejection


@integration
def test_malformed_timestamp_row_is_rejected_others_import(db_session, tmp_path):
    path = make_xlsx(
        tmp_path,
        COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 4.0],
            ["TL-1", "Haifa", "Dairy", "tomorrow-ish", 4.1],
            ["TL-1", "Haifa", "Dairy", None, 4.2],
            [None, None, None, None, None],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:30", 4.3],
            [None, "Haifa", "Dairy", "2026-09-14 06:45", 4.4],
        ],
    )
    result = ingest_xlsx(db_session, path)
    assert (result.rows_read, result.inserted, result.blank_rows) == (5, 2, 1)
    assert [(r.row_number, "tomorrow-ish" in r.reason) for r in result.rejected[:1]] == [(3, True)]
    assert [r.row_number for r in result.rejected] == [3, 4, 7]
    assert "missing" in result.rejected[1].reason
    assert result.rejected[2].reason == "missing logger"
    assert result.rows_read == result.inserted + result.duplicates + len(result.conflicts) + len(result.rejected)
    assert [r.source_row for r in readings(db_session)] == [2, 6]


@integration
def test_rejected_rows_create_no_entities(db_session, tmp_path):
    path = make_xlsx(tmp_path, COMBINED, [["TL-X", "Nowhere", "Void", "garbage", 4.0]])
    result = ingest_xlsx(db_session, path)
    assert len(result.rejected) == 1
    assert counts(db_session) == {"fridges": 0, "loggers": 0, "logger_assignments": 0, "temperature_readings": 0}


@integration
def test_in_batch_exact_duplicate_vs_conflict_first_wins(db_session, tmp_path):
    path = make_xlsx(
        tmp_path,
        COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 3.9],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 3.9],  # exact duplicate
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 7.7],  # conflict, first wins
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", "ERR"],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", "bad"],  # NULL vs NULL: duplicate
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", 5.0],  # NULL vs value: conflict
        ],
    )
    result = ingest_xlsx(db_session, path)
    assert (result.inserted, result.duplicates, len(result.conflicts)) == (2, 2, 2)
    c = result.conflicts[0]
    assert (c.row_number, c.logger_code, c.measured_at) == (4, "TL-1", datetime(2026, 9, 14, 6, 0))
    assert (c.kept, c.ignored, c.kept_source, c.kept_row_number) == (Decimal("3.90"), Decimal("7.70"), "batch", 2)
    assert (result.conflicts[1].kept, result.conflicts[1].ignored) == (None, Decimal("5.00"))
    assert [r.temperature_c for r in readings(db_session)] == [Decimal("3.90"), None]


@integration
def test_duplicates_and_conflicts_against_existing_database_rows(db_session, tmp_path):
    first = make_xlsx(
        tmp_path, COMBINED,
        [["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 3.9], ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", "ERR"]],
        "first.xlsx",
    )
    ingest_xlsx(db_session, first)
    second = make_xlsx(
        tmp_path, COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 3.9],  # identical -> duplicate
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 8.0],  # differs -> conflict
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", "ERR"],  # NULL == NULL -> duplicate
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:30", 4.4],  # new
        ],
        "second.xlsx",
    )
    result = ingest_xlsx(db_session, second)
    assert (result.inserted, result.duplicates, len(result.conflicts)) == (1, 2, 1)
    c = result.conflicts[0]
    assert (c.row_number, c.kept, c.ignored, c.kept_source, c.kept_row_number) == (
        3, Decimal("3.90"), Decimal("8.00"), "database", None,
    )
    assert [r.temperature_c for r in readings(db_session)] == [Decimal("3.90"), None, Decimal("4.40")]
    assert readings(db_session)[0].source_file == "first.xlsx"  # first wins, untouched


@integration
def test_reingest_inserts_nothing(db_session, tmp_path):
    path = make_xlsx(
        tmp_path, COMBINED,
        [["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 3.9], ["TL-2", "Tel Aviv", "Walk-in", "2026-09-14 06:00", "ERR"]],
    )
    ingest_xlsx(db_session, path)
    before = counts(db_session)
    again = ingest_xlsx(db_session, path)
    assert (again.inserted, again.duplicates, again.conflicts, again.rejected) == (0, 2, [], [])
    assert (again.loggers_created, again.fridges_created, again.assignments_created) == (0, 0, 0)
    assert counts(db_session) == before


@integration
def test_fahrenheit_logger_is_converted_and_existing_unit_is_not_overwritten(db_session, tmp_path):
    path = make_xlsx(
        tmp_path, COMBINED,
        [
            ["TL-F", "Haifa", "Dairy", "2026-09-14 06:00", 38.3],
            ["TL-F", "Haifa", "Dairy", "2026-09-14 06:15", 39],
            ["TL-F", "Haifa", "Dairy", "2026-09-14 06:30", 32],
            ["TL-F", "Haifa", "Dairy", "2026-09-14 06:45", 212],
            ["TL-C", "Haifa", "Dairy", "2026-09-14 06:00", 38.3],  # not in mapping: Celsius
        ],
    )
    ingest_xlsx(db_session, path, logger_units={"TL-F": "F"})
    values = {(r.logger_id, r.raw_value): r.temperature_c for r in readings(db_session)}
    tl_f = db_session.scalar(select(Logger).where(Logger.code == "TL-F"))
    tl_c = db_session.scalar(select(Logger).where(Logger.code == "TL-C"))
    assert [values[(tl_f.id, k)] for k in ("38.3", "39", "32", "212")] == [
        Decimal("3.50"), Decimal("3.89"), Decimal("0.00"), Decimal("100.00"),
    ]
    assert values[(tl_c.id, "38.3")] == Decimal("38.30")
    assert (tl_f.unit, tl_c.unit) == ("F", "C")

    # A later mapping never changes an existing logger's unit.
    path2 = make_xlsx(tmp_path, COMBINED, [["TL-F", "Haifa", "Dairy", "2026-09-15 06:00", 40]], "again.xlsx")
    ingest_xlsx(db_session, path2, logger_units={"TL-F": "C"})
    db_session.refresh(tl_f)
    assert tl_f.unit == "F"


@integration
def test_source_units_convert_without_changing_native_unit(db_session, tmp_path):
    create = make_xlsx(
        tmp_path, COMBINED,
        [["TL-0231", "Haifa", "Dairy", "2026-09-14 06:00", 38.0]],
        "native.xlsx",
    )
    ingest_xlsx(db_session, create, logger_units={"TL-0231": "F"})
    logger = db_session.scalar(select(Logger).where(Logger.code == "TL-0231"))
    assert logger.unit == "F"

    already_c = make_xlsx(
        tmp_path, COMBINED,
        [["TL-0231", "Haifa", "Dairy", "2026-09-14 07:00", 3.8]],
        "celsius.xlsx",
    )
    ingest_xlsx(db_session, already_c, source_units={"TL-0231": "C"})
    db_session.refresh(logger)
    assert logger.unit == "F"
    later = db_session.scalars(
        select(TemperatureReading).where(
            TemperatureReading.measured_at == datetime(2026, 9, 14, 7, 0)
        )
    ).one()
    assert later.temperature_c == Decimal("3.80")


@integration
def test_assignments_are_undated_and_none_invented(db_session, tmp_path):
    path = make_xlsx(
        tmp_path, COMBINED,
        [
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 4.0],
            ["TL-1", "Haifa", "Dairy", "2026-09-14 06:15", 4.0],
            ["TL-2", "Tel Aviv", "Walk-in", "2026-09-14 06:00", 4.0],
        ],
    )
    ingest_xlsx(db_session, path)
    ingest_xlsx(db_session, path)
    assignments = db_session.scalars(select(LoggerAssignment)).all()
    assert len(assignments) == 2
    assert all(a.assigned_from is None and a.assigned_to is None for a in assignments)


@integration
def test_existing_dated_assignment_is_left_alone_and_no_undated_one_added(db_session, tmp_path):
    logger = Logger(code="TL-1", unit="C")
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add_all([logger, fridge])
    db_session.flush()
    db_session.add(LoggerAssignment(logger_id=logger.id, fridge_id=fridge.id, assigned_from=datetime(2026, 9, 1)))
    db_session.flush()
    path = make_xlsx(tmp_path, COMBINED, [["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 4.0]])
    result = ingest_xlsx(db_session, path)
    assert result.assignments_created == 0
    (assignment,) = db_session.scalars(select(LoggerAssignment)).all()
    assert assignment.assigned_from == datetime(2026, 9, 1) and assignment.assigned_to is None


@integration
def test_logger_under_two_fridges_keeps_per_reading_fridge(db_session, tmp_path):
    path = make_xlsx(
        tmp_path, COMBINED,
        [
            ["TL-0417", "Tel Aviv", "Walk-in", "2026-09-14 06:00", 4.1],
            ["TL-0417", "tel aviv", "Display 2", "2026-09-17 06:00", 3.7],
            ["TL-0417", "TEL  AVIV", "walk-in", "2026-09-14 06:15", 9.4],  # same fridge, other spelling
        ],
    )
    result = ingest_xlsx(db_session, path)
    assert (result.fridges_created, result.assignments_created, result.inserted) == (2, 2, 3)
    fridges = {f.name: f for f in db_session.scalars(select(Fridge))}
    assert set(fridges) == {"Walk-in", "Display 2"}
    assert {f.branch_name for f in fridges.values()} == {"Tel Aviv"}  # first-seen spelling kept
    by_time = {r.measured_at: r.fridge_id for r in readings(db_session)}
    assert by_time[datetime(2026, 9, 14, 6, 0)] == fridges["Walk-in"].id
    assert by_time[datetime(2026, 9, 14, 6, 15)] == fridges["Walk-in"].id
    assert by_time[datetime(2026, 9, 17, 6, 0)] == fridges["Display 2"].id
    assert len(db_session.scalars(select(LoggerAssignment)).all()) == 2


@integration
def test_names_match_existing_rows_case_insensitively_keeping_stored_spelling(db_session, tmp_path):
    db_session.add_all([Fridge(branch_name="Tel Aviv", name="Walk-in"), Logger(code="TL-0417", unit="C")])
    db_session.flush()
    path = make_xlsx(tmp_path, COMBINED, [["tl-0417", "  tel   aviv ", "WALK-IN", "2026-09-14 06:00", 4.0]])
    result = ingest_xlsx(db_session, path)
    assert (result.loggers_created, result.fridges_created, result.inserted) == (0, 0, 1)
    assert db_session.scalar(select(func.count()).select_from(Fridge)) == 1
    assert db_session.scalar(select(Fridge.branch_name)) == "Tel Aviv"


@integration
def test_service_does_not_commit(db_session, tmp_path):
    path = make_xlsx(tmp_path, COMBINED, [["TL-1", "Haifa", "Dairy", "2026-09-14 06:00", 4.0]])
    commits = []
    from sqlalchemy import event

    event.listen(db_session, "after_commit", lambda s: commits.append(1))
    ingest_xlsx(db_session, path)
    assert commits == []
    assert db_session.in_transaction()
    assert counts(db_session)["temperature_readings"] == 1


@integration
def test_ingest_leaves_nothing_behind_after_rollback(db_engine, tmp_path):
    from sqlalchemy.orm import Session

    path = make_xlsx(tmp_path, COMBINED, [["TL-ROLLBACK", "Haifa", "Dairy", "2026-09-14 06:00", 4.0]])
    with Session(db_engine) as session:
        ingest_xlsx(session, path)
        session.rollback()  # the caller owns the transaction
    with Session(db_engine) as session:
        assert session.scalar(select(func.count()).select_from(Logger).where(Logger.code == "TL-ROLLBACK")) == 0
        assert session.scalar(select(func.count()).select_from(TemperatureReading)) == 0


@integration
def test_header_temp_unit_matching_logger_is_imported(db_session, tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add(Logger(code="TL-F", unit="F"))
    db_session.add(fridge)
    db_session.flush()
    path = make_xlsx(tmp_path, ["Time", "Temp F"], [["2026-09-14 06:00", 39]])
    result = ingest_xlsx(db_session, path, logger="TL-F", fridge=fridge)
    assert result.inserted == 1
    assert readings(db_session)[0].temperature_c == Decimal("3.89")


@integration
def test_header_temp_unit_conflicting_with_logger_is_rejected(db_session, tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add(Logger(code="TL-C", unit="C"))
    db_session.add(fridge)
    db_session.flush()
    path = make_xlsx(tmp_path, ["Time", "Temperature (°F)"], [["2026-09-14 06:00", 39]])
    with pytest.raises(IngestError, match="indicates F"):
        ingest_xlsx(db_session, path, logger="TL-C", fridge=fridge)
    assert counts(db_session)["temperature_readings"] == 0


@integration
def test_celsius_header_conflicts_with_fahrenheit_logger(db_session, tmp_path):
    fridge = Fridge(branch_name="Haifa", name="Dairy")
    db_session.add(Logger(code="TL-F2", unit="F"))
    db_session.add(fridge)
    db_session.flush()
    path = make_xlsx(tmp_path, ["Time", "Temp C"], [["2026-09-14 06:00", 4.0]])
    with pytest.raises(IngestError, match="indicates C"):
        ingest_xlsx(db_session, path, logger="TL-F2", fridge=fridge)
